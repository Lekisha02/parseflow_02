"""Generate Office/text/image test documents with known content (python-docx, python-pptx, openpyxl, Pillow)."""
import io

from PIL import Image, ImageDraw, ImageFont

TABLE = [["Algorithm", "Best", "Worst"], ["Insertion sort", "n", "n^2"], ["Merge sort", "n log n", "n log n"]]
PARA = ("Insertion sort builds the sorted array one element at a time. Each new key is compared with the elements "
        "to its left and inserted where it belongs.")


def png_bytes(text="FIG", size=(240, 120), color=(200, 220, 255)):
    im = Image.new("RGB", size, color)
    d = ImageDraw.Draw(im)
    d.rectangle([4, 4, size[0] - 5, size[1] - 5], outline=(0, 0, 0), width=3)
    try:
        f = ImageFont.load_default(size=36)
    except TypeError:
        f = ImageFont.load_default()
    d.text((20, 30), text, fill=(0, 0, 0), font=f)
    buf = io.BytesIO(); im.save(buf, "PNG")
    return buf.getvalue()


def make_docx(path):
    import docx
    from docx.shared import Inches
    d = docx.Document()
    d.add_heading("Sorting Algorithms", level=1)
    d.add_paragraph(PARA)
    d.add_heading("Comparison", level=2)
    t = d.add_table(rows=len(TABLE) + 1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].merge(hdr[2]).text = "Running times"
    for r, row in enumerate(TABLE, start=1):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_paragraph("Stable sorts", style="List Bullet")
    d.add_paragraph("In-place sorts", style="List Bullet")
    p = d.add_paragraph("The running time is O(n")
    r = p.add_run("2"); r.font.superscript = True
    p.add_run(") in the worst case.")
    d.add_picture(io.BytesIO(png_bytes("DOCX")), width=Inches(2))
    d.save(path)
    return path


def make_pptx(path):
    import pptx
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Graph Algorithms"
    s.placeholders[1].text = "Lecture 7"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Breadth-first search"
    body = s.placeholders[1].text_frame
    body.text = "Visits vertices in order of distance"
    p = body.add_paragraph(); p.text = "Uses a FIFO queue"; p.level = 1
    s.notes_slide.notes_text_frame.text = "Remind students that BFS gives shortest paths in unweighted graphs."
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Running times"
    tb = s.shapes.add_table(len(TABLE), 3, Inches(1), Inches(2), Inches(8), Inches(2)).table
    for r, row in enumerate(TABLE):
        for c, v in enumerate(row):
            tb.cell(r, c).text = v
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Edges per graph"
    cd = CategoryChartData(); cd.categories = ["G1", "G2", "G3"]; cd.add_series("Edges", (12, 30, 7))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.8), Inches(6), Inches(4), cd)
    s.shapes.add_picture(io.BytesIO(png_bytes("PPTX")), Inches(7.2), Inches(5.6), width=Inches(2))
    prs.save(path)
    return path


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Sales"
    ws.append(["Quarterly sales", None, None]); ws.merge_cells("A1:C1")
    ws.append(["Region", "Q1", "Q2"])
    ws.append(["North", 10, 12.5]); ws.append(["South", 7, 9]); ws.append(["Total", "=SUM(B3:B4)", "=SUM(C3:C4)"])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Item", "Value"]); ws2.append(["Growth", 0.25]); ws2["B2"].number_format = "0%"
    wb.save(path)
    return path


def make_csv(path):
    with open(path, "w", newline="") as f:
        f.write("name;score;grade\nAsha;91;A\nRavi;78;B\n")
    return path


def make_md(path):
    with open(path, "w") as f:
        f.write("# Heaps\n\nA **max-heap** keeps every parent at least as large as its children.\n\n"
                "- insert: O(log n)\n- extract-max: O(log n)\n\n| Op | Cost |\n|---|---|\n| build | O(n) |\n\n```\nMAX-HEAPIFY(A, i)\n```\n")
    return path


def make_image(path, fmt, frames=1, exif_rotate=False):
    ims = []
    for k in range(frames):
        im = Image.new("RGB", (900, 400), "white")
        d = ImageDraw.Draw(im)
        try:
            f = ImageFont.truetype("DejaVuSans.ttf", 44)
        except Exception:
            try:
                f = ImageFont.load_default(size=44)
            except TypeError:
                f = ImageFont.load_default()
        d.text((40, 60), f"Page {k + 1} scanned text", fill="black", font=f)
        d.text((40, 160), "Binary search halves the range", fill="black", font=f)
        ims.append(im)
    if exif_rotate:                         # stored sideways with an EXIF 'rotate 90' tag, as phones do
        im = ims[0].rotate(90, expand=True)
        ex = im.getexif(); ex[0x0112] = 6
        im.save(path, "JPEG", exif=ex.tobytes())
        return path
    if frames > 1:
        ims[0].save(path, fmt, save_all=True, append_images=ims[1:], dpi=(150, 150))
    else:
        ims[0].save(path, fmt)
    return path
